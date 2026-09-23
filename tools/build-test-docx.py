from docx import Document
from docx.shared import Mm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT, WD_ROW_HEIGHT_RULE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from pathlib import Path

ROOT=Path(__file__).resolve().parents[1]
OUT=ROOT/'worksheets'/'past_tense_test_A_B.docx'
OUT.parent.mkdir(exist_ok=True)

A=[
('played','I ________________ basketball.','나는 농구를 했습니다.','play / played'),
('hugged','I ________________ my grandma.','나는 할머니를 껴안았습니다.','hug / hugged'),
('picked','I ________________ watermelons.','나는 수박을 땄습니다.','pick / picked'),
('baked','I ________________ cookies.','나는 쿠키를 구웠습니다.','bake / baked'),
('swam','I ________________ in the sea.','나는 바다에서 수영했습니다.','swimmed / swam'),
('bought','I ________________ fruit.','나는 과일을 샀습니다.','buyed / bought'),
('ran','I ________________ fast.','나는 빨리 달렸습니다.','runned / ran'),
('sat','I ________________ on the chair.','나는 의자에 앉았습니다.','sitted / sat'),
('drank','I ________________ milk.','나는 우유를 마셨습니다.','drinked / drank'),
('read','I ________________ a book.','나는 책을 읽었습니다.',None),
('wrote','I ________________ a letter.','나는 편지를 썼습니다.',None),
('came','I ________________ to school.','나는 학교에 왔습니다.',None)]
B=[
('walked','I ________________ in the park.','나는 공원에서 걸었습니다.','walk / walked'),
('talked','I ________________ with my friend.','나는 친구와 이야기했습니다.','talk / talked'),
('washed','I ________________ my hands.','나는 손을 씻었습니다.','wash / washed'),
('opened','I ________________ the door.','나는 문을 열었습니다.','open / opened'),
('ate','I ________________ lunch.','나는 점심을 먹었습니다.','eated / ate'),
('listened','I ________________ to music.','나는 음악을 들었습니다.','listenned / listened'),
('danced','I ________________.','나는 춤을 췄습니다.','danceed / danced'),
('studied','I ________________ English.','나는 영어를 공부했습니다.','studyed / studied'),
('stopped','I ________________ the bicycle.','나는 자전거를 멈췄습니다.','stoped / stopped'),
('bought','I ________________ fruit.','나는 과일을 샀습니다.',None),
('swam','I ________________ in the sea.','나는 바다에서 수영했습니다.',None),
('read','I ________________ a book.','나는 책을 읽었습니다.',None)]

def img(name):
    for p in [ROOT/'img'/f'{name}.png',ROOT/'img'/'supplementary'/f'{name}.webp']:
        if p.exists(): return str(p)
    raise FileNotFoundError(name)

def noborders(t):
    pr=t._tbl.tblPr; b=OxmlElement('w:tblBorders')
    for e in ['top','left','bottom','right','insideH','insideV']:
        x=OxmlElement('w:'+e); x.set(qn('w:val'),'nil'); b.append(x)
    pr.append(b)

doc=Document(); sec=doc.sections[0]
sec.page_width=Mm(210); sec.page_height=Mm(297); sec.top_margin=Mm(7); sec.bottom_margin=Mm(6); sec.left_margin=Mm(10); sec.right_margin=Mm(10)
doc.styles['Normal'].font.name='Arial'; doc.styles['Normal']._element.rPr.rFonts.set(qn('w:eastAsia'),'Malgun Gothic')

def page(label,qs):
    p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_after=Pt(1)
    r=p.add_run(f'Unit 7  Past Tense Check · Test {label}'); r.bold=True; r.font.size=Pt(16); r.font.color.rgb=RGBColor(31,78,121)
    p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_after=Pt(2)
    r=p.add_run('그림을 보고 빈칸에 알맞은 과거형을 쓰세요.'); r.bold=True; r.font.size=Pt(10.5)
    p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.RIGHT; p.paragraph_format.space_after=Pt(2); p.add_run('5학년 ____반   이름 ____________________')
    t=doc.add_table(rows=6,cols=2); t.alignment=WD_TABLE_ALIGNMENT.CENTER
    for i,q in enumerate(qs):
        c=t.cell(i//2,i%2); c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
        inn=c.add_table(rows=1,cols=2); noborders(inn)
        ip=inn.cell(0,0).paragraphs[0]; ip.alignment=WD_ALIGN_PARAGRAPH.CENTER
        ip.add_run().add_picture(img(q[0]),width=Mm(35))
        tc=inn.cell(0,1)
        p=tc.paragraphs[0]; rr=p.add_run(f'{i+1}. {q[1]}'); rr.bold=True; rr.font.size=Pt(12.5)
        p=tc.add_paragraph(); rr=p.add_run(q[2]); rr.font.size=Pt(9); rr.font.color.rgb=RGBColor(90,90,90)
        p=tc.add_paragraph()
        if q[3]:
            rr=p.add_run(q[3]); rr.bold=True; rr.font.size=Pt(10); rr.font.color.rgb=RGBColor(0,112,84)
        else:
            rr=p.add_run('보기 없이 직접 쓰기'); rr.bold=True; rr.font.size=Pt(9); rr.font.color.rgb=RGBColor(192,80,77)
    for row in t.rows: row.height=Mm(39); row.height_rule=WD_ROW_HEIGHT_RULE.AT_LEAST

page('A',A); doc.add_page_break(); page('B',B)
doc.save(OUT)
print(OUT)
